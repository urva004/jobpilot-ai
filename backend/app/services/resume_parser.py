import io
import pdfplumber
import docx

# Minimum meaningful text length from pdfplumber before triggering OCR fallback
_OCR_THRESHOLD = 50

class ResumeParser:
    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        """Extract text from PDF. Falls back to OCR if pdfplumber yields < 50 chars (scanned PDF)."""
        text = ""
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            print(f"[ResumeParser] pdfplumber failed: {e}")

        if len(text.strip()) >= _OCR_THRESHOLD:
            return text

        # ── Feature 4: OCR Fallback ────────────────────────────────────────────
        print(f"[ResumeParser] pdfplumber extracted {len(text.strip())} chars — triggering OCR fallback")
        try:
            import pytesseract
            from pdf2image import convert_from_bytes

            images = convert_from_bytes(file_bytes, dpi=200)
            ocr_text = ""
            for i, img in enumerate(images):
                page_text = pytesseract.image_to_string(img)
                if page_text:
                    ocr_text += page_text + "\n"
                    print(f"[ResumeParser] OCR page {i+1}: {len(page_text)} chars extracted")

            if len(ocr_text.strip()) >= _OCR_THRESHOLD:
                print(f"[ResumeParser] OCR succeeded: {len(ocr_text.strip())} total chars")
                return ocr_text
            else:
                print(f"[ResumeParser] OCR also yielded sparse text ({len(ocr_text.strip())} chars). Returning best available.")
                return ocr_text or text

        except ImportError as ie:
            missing = "pytesseract" if "pytesseract" in str(ie) else "pdf2image"
            print(f"[ResumeParser] OCR library '{missing}' not installed — cannot fall back to OCR.")
            raise ValueError(
                f"This PDF appears to be a scanned image with no extractable text. "
                f"OCR fallback requires '{missing}' to be installed. "
                f"Please convert your PDF to text-based format or install OCR dependencies."
            )
        except Exception as e:
            print(f"[ResumeParser] OCR fallback failed: {e}")
            # Return whatever pdfplumber gave us rather than crashing the upload
            return text

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        doc = docx.Document(io.BytesIO(file_bytes))
        full_text = []
        for para in doc.paragraphs:
            if para.text:
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                # Deduplicate row text to avoid repeating identical text across merged cells
                row_text = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and (not row_text or row_text[-1] != cell_text):
                        row_text.append(cell_text)
                if row_text:
                    full_text.append(" | ".join(row_text))
        return "\n".join(full_text)

    @classmethod
    def parse(cls, filename: str, file_bytes: bytes) -> str:
        ext = filename.split(".")[-1].lower()
        if ext == "pdf":
            return cls.extract_text_from_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            return cls.extract_text_from_docx(file_bytes)
        else:
            raise ValueError(f"Unsupported file format: {ext}. Only PDF and DOCX are supported.")
